"""Fill the Test Execution sheet of the Excel report from Allure results.

Usage:
    python tools/fill_excel_report.py
    python tools/fill_excel_report.py --results reports/allure-results --executed-by Nour
"""
import argparse
import glob
import json
import os
import re
from datetime import datetime

from openpyxl import load_workbook

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_RESULTS = os.path.join(ROOT_DIR, "reports", "allure-results")
DEFAULT_TEMPLATE = os.path.join(ROOT_DIR, "docs", "SauceDemo_Automation_Test_Report.xlsx")
DEFAULT_OUTPUT = os.path.join(ROOT_DIR, "reports", "SauceDemo_Automation_Test_Report_Filled.xlsx")
REPORT_URL = "https://nourn117.github.io/saucedemo-selenium-automation/"

ID_PATTERN = re.compile(r"(AUT-\d{3})")

PASSED = "Passed"
FAILED = "Failed"
XFAIL = "Known Defect (XFail)"
XPASS = "Unexpected Pass (XPass)"
BLOCKED = "Blocked"
SKIPPED = "Skipped"

# Worst status wins when one test has several results (parametrized tests)
SEVERITY = [FAILED, XPASS, BLOCKED, XFAIL, SKIPPED, PASSED]

EXEC_COLUMNS = {"id": 1, "date": 7, "status": 8, "actual": 9, "duration": 10, "evidence": 11, "remarks": 13}
DEFECT_COLUMNS = {"automation_tc": 3, "reproduced": 13}


def map_status(result):
    status = result.get("status", "")
    message = (result.get("statusDetails") or {}).get("message", "") or ""
    if "XPASS" in message.upper():
        return XPASS
    if status == "passed":
        return PASSED
    if status == "skipped":
        return XFAIL if "XFAIL" in message.upper() else SKIPPED
    if status == "failed":
        return FAILED
    if status == "broken":
        return BLOCKED if "WebDriverException" in message else FAILED
    return SKIPPED


def first_line(text, limit=300):
    text = (text or "").strip().splitlines()
    return text[0][:limit] if text else ""


def load_results(results_dir):
    latest = {}
    for path in glob.glob(os.path.join(results_dir, "*-result.json")):
        with open(path, encoding="utf-8") as file:
            result = json.load(file)
        match = ID_PATTERN.search(result.get("name", "")) or ID_PATTERN.search(result.get("fullName", ""))
        if not match:
            continue
        key = result.get("historyId") or result.get("uuid")
        previous = latest.get(key)
        if previous is None or result.get("stop", 0) > previous["result"].get("stop", 0):
            latest[key] = {"id": match.group(1), "result": result}

    grouped = {}
    for item in latest.values():
        grouped.setdefault(item["id"], []).append(item["result"])
    return grouped


def summarize(results):
    statuses = [map_status(r) for r in results]
    status = min(statuses, key=SEVERITY.index)
    duration = sum(max(r.get("stop", 0) - r.get("start", 0), 0) for r in results) / 1000
    stop = max(r.get("stop", 0) for r in results)

    if status == PASSED:
        actual = "As expected"
    else:
        worst = next(r for r, s in zip(results, statuses) if s == status)
        message = first_line((worst.get("statusDetails") or {}).get("message"))
        prefix = {XFAIL: "Known defect reproduced", XPASS: "Passed although marked as known defect"}.get(status, "")
        actual = f"{prefix}: {message}" if prefix and message else (prefix or message or status)

    remarks = f"{len(results)} parametrized runs" if len(results) > 1 else None
    return {"status": status, "actual": actual, "duration": round(duration, 2),
            "date": datetime.fromtimestamp(stop / 1000) if stop else None, "remarks": remarks}


def read_run_number(results_dir):
    path = os.path.join(results_dir, "executor.json")
    if os.path.exists(path):
        with open(path, encoding="utf-8") as file:
            return json.load(file).get("buildName")
    return "Local run"


def set_summary_field(sheet, label, value):
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == label:
            sheet.cell(row=row, column=2).value = value
            return


def fill(template, results_dir, output, executed_by):
    grouped = load_results(results_dir)
    if not grouped:
        raise SystemExit(f"No results with AUT-xxx IDs found in {results_dir}")

    workbook = load_workbook(template)
    execution = workbook["Test Execution"]
    statuses = {}
    filled = 0

    for row in range(2, execution.max_row + 1):
        test_id = execution.cell(row=row, column=EXEC_COLUMNS["id"]).value
        if not test_id:
            continue
        if test_id not in grouped:
            execution.cell(row=row, column=EXEC_COLUMNS["status"]).value = "Not Run"
            continue
        summary = summarize(grouped[test_id])
        statuses[test_id] = summary["status"]
        execution.cell(row=row, column=EXEC_COLUMNS["date"]).value = summary["date"]
        execution.cell(row=row, column=EXEC_COLUMNS["status"]).value = summary["status"]
        execution.cell(row=row, column=EXEC_COLUMNS["actual"]).value = summary["actual"]
        execution.cell(row=row, column=EXEC_COLUMNS["duration"]).value = summary["duration"]
        execution.cell(row=row, column=EXEC_COLUMNS["evidence"]).value = f"{REPORT_URL} (search {test_id})"
        if summary["remarks"]:
            execution.cell(row=row, column=EXEC_COLUMNS["remarks"]).value = summary["remarks"]
        filled += 1

    defects = workbook["Defect Report"]
    for row in range(2, defects.max_row + 1):
        linked = defects.cell(row=row, column=DEFECT_COLUMNS["automation_tc"]).value or ""
        linked_statuses = [statuses[i] for i in ID_PATTERN.findall(linked) if i in statuses]
        if not linked_statuses:
            continue
        reproduced = any(s in (XFAIL, FAILED) for s in linked_statuses)
        defects.cell(row=row, column=DEFECT_COLUMNS["reproduced"]).value = "Yes" if reproduced else "No"

    summary_sheet = workbook["Summary"]
    dates = [execution.cell(row=r, column=EXEC_COLUMNS["date"]).value for r in range(2, execution.max_row + 1)]
    dates = [d for d in dates if d]
    if dates:
        set_summary_field(summary_sheet, "Execution Date", max(dates).strftime("%d/%m/%Y"))
    set_summary_field(summary_sheet, "Run / Build No.", read_run_number(results_dir))
    if executed_by:
        set_summary_field(summary_sheet, "Executed By", executed_by)

    workbook.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(output), exist_ok=True)
    workbook.save(output)
    print(f"Filled {filled} test cases -> {output}")


def main():
    parser = argparse.ArgumentParser(description="Fill the Excel test report from Allure results")
    parser.add_argument("--results", default=DEFAULT_RESULTS)
    parser.add_argument("--template", default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--executed-by", default="")
    args = parser.parse_args()
    fill(args.template, args.results, args.output, args.executed_by)


if __name__ == "__main__":
    main()
